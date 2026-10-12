"""User-submitted sources share the same evidence/routing boundary as channels."""
import base64
import hashlib
import io
import json
from pathlib import Path
from typing import Any, Dict, Optional

from fastapi import HTTPException
from sqlalchemy.orm import Session

from app.models.project import Project
from app.services.source_intelligence_pipeline import SourceIntelligencePipeline

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARS = 200_000


def validate_target(project_id: Optional[str], db: Session) -> None:
    if project_id and not db.query(Project).filter(
        Project.id == project_id, Project.is_system.is_(False)
    ).first():
        raise HTTPException(404, "Target project does not exist or is a system project.")


def process_input(source: str, payload: Dict[str, Any], db: Session, tenant_id: str,
                  actor_id: str, project_id: Optional[str] = None,
                  source_event_id: Optional[str] = None) -> Dict[str, Any]:
    validate_target(project_id, db)
    text = str(payload.get("text", "")).strip()
    if not text:
        raise HTTPException(400, "Input must contain readable text.")
    if len(text) > MAX_TEXT_CHARS:
        raise HTTPException(413, "Extracted text exceeds 200,000 characters.")
    # Content-addressed within actor + tenant + target: safe retry, distinct routing.
    digest = hashlib.sha256(json.dumps(
        [tenant_id, actor_id, project_id,
         {"filename": payload.get("filename"), "sha256": payload["sha256"]} if payload.get("sha256") else payload],
        sort_keys=True, default=str
    ).encode()).hexdigest()
    outcome = SourceIntelligencePipeline().process(
        source=source, payload=payload, db=db, tenant_id=tenant_id,
        actor_id=actor_id, project_id=project_id,
        source_event_id=source_event_id or f"{source}_{digest}",
        event_type="image_upload" if source == "image_upload" else "manual_submission",
        metadata={"filename": payload.get("filename"), "submitted_by": actor_id},
        visual_context=text if source == "image_upload" else None,
        ignore_unresolved=source in {"file_upload", "image_upload"},
    )
    return outcome.model_dump()


def extract_upload(filename: str, data: bytes) -> tuple[str, Dict[str, Any]]:
    if not data:
        raise HTTPException(400, "The uploaded file is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, "Uploads are limited to 10 MB.")
    suffix = Path(filename).suffix.lower()
    payload: Dict[str, Any] = {"filename": Path(filename).name,
                              "sha256": hashlib.sha256(data).hexdigest()}
    source = "file_upload"
    try:
        if suffix in {".txt", ".md", ".csv", ".tsv", ".vtt", ".srt"}:
            text = data.decode("utf-8-sig")
        elif suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ValueError("Encrypted PDFs are not supported")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif suffix == ".docx":
            from docx import Document
            document = Document(io.BytesIO(data))
            text = "\n".join([p.text for p in document.paragraphs] + [
                "\t".join(cell.text for cell in row.cells)
                for table in document.tables for row in table.rows
            ])
        elif suffix in {".png", ".jpg", ".jpeg", ".webp"}:
            from PIL import Image, UnidentifiedImageError
            from app.services.multimodal_service import MultimodalService
            try:
                with Image.open(io.BytesIO(data)) as image:
                    image.verify()
                    image_format = image.format
                expected = {".png": "PNG", ".jpg": "JPEG", ".jpeg": "JPEG", ".webp": "WEBP"}[suffix]
                if image_format != expected:
                    raise ValueError("Image content does not match its extension")
            except (UnidentifiedImageError, OSError, ValueError) as exc:
                raise HTTPException(400, "Invalid or mismatched image file.") from exc
            mime = {"PNG": "image/png", "JPEG": "image/jpeg", "WEBP": "image/webp"}[image_format]
            analysis = MultimodalService().extract_from_image(data, mime_type=mime)
            if not analysis.get("processed"):
                raise HTTPException(503, "Image analysis unavailable. Configure NVIDIA_API_KEY and retry.")
            lines = [str(analysis.get(key) or "") for key in ("ocr_text", "summary")]
            for key in ("decisions", "requirements"):
                lines.extend(f"{key[:-1].capitalize()}: {item}" for item in analysis.get(key, []) if item)
            for relation in analysis.get("relationships", []):
                if isinstance(relation, dict):
                    lines.append(f"Connection: {relation.get('source', '')} -> {relation.get('target', '')}: {relation.get('label', '')}")
            text = "\n".join(lines)
            payload.update(vision_analysis=analysis, mime_type=mime,
                           image_base64=base64.b64encode(data).decode("ascii"))
            source = "image_upload"
        else:
            raise HTTPException(415, "Supported files: TXT, Markdown, CSV, TSV, VTT, SRT, PDF, DOCX, PNG, JPEG and WebP.")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(400, f"Could not read file: {exc}") from exc
    payload["text"] = text.strip()
    if not payload["text"]:
        raise HTTPException(400, "No readable text found. For scanned documents, upload the page as an image.")
    return source, payload
